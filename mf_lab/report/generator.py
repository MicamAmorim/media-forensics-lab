from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt

from mf_lab.utils.io import read_json

ROOT = Path(__file__).resolve().parents[2]

METHOD_LABELS = {
    "hash_sha256": "Hash criptográfico SHA-256",
    "perceptual_hashes": "Hashes perceptuais (similaridade)",
    "c2pa": "Proveniência C2PA / Content Credentials",
    "metadata": "Metadados e estrutura técnica",
    "ela": "Error Level Analysis (ELA) - triagem",
    "histogram": "Histograma RGB - triagem",
    "noise_residual": "Residual de ruído - triagem",
    "noise_map": "Consistência local de ruído - triagem",
    "jpeg_ghost": "JPEG Ghost - triagem",
    "jpeg_quantization": "Tabelas de quantização JPEG",
    "jpeg_dct": "Periodicidade DCT / dupla compressão - heurística",
    "copy_move_orb": "Copy-move por características ORB - triagem",
    "frequency": "Análise em frequência FFT - triagem",
    "resampling": "Traços de reamostragem - triagem",
    "steganography_lsb": "Esteganálise LSB - triagem",
    "prnu_screen": "Residual tipo PRNU - triagem (não atribuição de câmera)",
    "face_artifacts": "Métricas de região facial - triagem",
    "face_context_consistency": "Consistência face versus contexto - triagem",
    "synthetic_spectral": "Triagem espectral de mídia sintética",
    "synthetic_feature_bank": "Banco de características de mídia sintética",
    "synthetic_ml": "Classificador ML handcrafted de mídia sintética",
    "synthetic_deep": "Detector profundo ONNX de mídia sintética",
    "synthetic_evidence_fusion": "Convergência de famílias de sinais sintéticos",
    "deepfake_protocol": "Protocolo MFLAB-DF para imagem",
    "video_timing": "Estrutura temporal / timestamps",
    "video_duplicates": "Frames adjacentes quase duplicados",
    "video_transition_anomalies": "Transições visuais abruptas - triagem",
    "video_motion_discontinuities": "Descontinuidades de movimento por fluxo óptico - triagem",
    "reference_image_difference": "Comparação de imagem com referência",
    "reference_video_alignment": "Alinhamento de vídeo com referência",
    "video_deepfake_protocol": "Protocolo MFLAB-DF para vídeo",
    "veritas_upstream_crosscheck": "Cross-check opcional CodeRafay/Veritas",
}


def _cant_split_row(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    if trPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit") is None:
        trPr.append(OxmlElement("w:cantSplit"))


def _repeat_header_row(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "true")
    trPr.append(tblHeader)


def _load_case(case_dir: Path):
    p = case_dir / "case.yaml"
    if not p.exists():
        tpl = yaml.safe_load((ROOT / "templates/case.yaml").read_text(encoding="utf-8"))
        tpl["case"]["id"] = case_dir.name
        p.write_text(yaml.safe_dump(tpl, allow_unicode=True, sort_keys=False), encoding="utf-8")
    return yaml.safe_load(p.read_text(encoding="utf-8"))["case"]


def _refs_for_reports(reports):
    refs = yaml.safe_load((ROOT / "bibliography/references.yaml").read_text(encoding="utf-8"))["references"]
    keys = {"cpp_art473", "cpp_chain_custody", "swgde_image_auth", "swgde_video_auth"}
    for r in reports:
        for m in r.get("method_registry", {}).values():
            keys.update(m.get("refs", []))
    return [refs[k] for k in sorted(keys) if k in refs]


def _deepfake_summary(report: dict) -> str | None:
    methods = report.get("methods", {})
    proto = methods.get("deepfake_protocol") or methods.get("video_deepfake_protocol")
    if not isinstance(proto, dict):
        return None
    version = proto.get("protocol_version", "MFLAB-DF")
    triage = proto.get("triage_assessment", "não informado")
    validated = proto.get("validated_external_models", 0)
    families = proto.get("evidence_families", []) or []
    observations = proto.get("screening_observations", []) or []
    conclusion = proto.get("evidentiary_conclusion", "inconclusivo")
    return (
        f"{version}: triagem={triage}; observações heurísticas não calibradas={len(observations)}; "
        f"famílias de evidência validadas/sinalizadas={len(families)}; modelos externos explicitamente "
        f"marcados como validados={validated}; conclusão automática de valor probatório={conclusion}."
    )


def _assessment(reports):
    if not reports:
        return "Não foram encontrados relatórios técnicos para consolidação."
    flags = []
    for r in reports:
        m = r.get("methods", {})
        name = Path(r.get("file", "arquivo")).name
        cm = m.get("copy_move_orb", {}).get("score", 0)
        gaps = len(m.get("video_timing", {}).get("large_gaps", []))
        dups = m.get("video_duplicates", {}).get("duplicate_count", 0)
        abrupt = m.get("video_transition_anomalies", {}).get("anomaly_count", 0)
        motion = m.get("video_motion_discontinuities", {}).get("anomaly_count", 0)
        proto = m.get("deepfake_protocol") or m.get("video_deepfake_protocol") or {}
        if cm >= 0.5:
            flags.append(f"{name}: elevada quantidade de correspondências ORB compatíveis com copy-move em triagem.")
        if gaps:
            flags.append(f"{name}: {gaps} descontinuidade(s) temporal(is) potencial(is) segundo timestamps.")
        if dups:
            flags.append(f"{name}: {dups} transição(ões) entre frames atingiram o limiar de quase-duplicação.")
        if abrupt:
            idx = [x.get("index") for x in m.get("video_transition_anomalies", {}).get("anomalous_transitions", [])]
            flags.append(f"{name}: {abrupt} transição(ões) abrupta(s) foram sinalizadas nos índices {idx}.")
        if motion:
            idx = [x.get("index") for x in m.get("video_motion_discontinuities", {}).get("anomalies", [])]
            flags.append(f"{name}: {motion} descontinuidade(s) de movimento foram sinalizadas nos índices {idx}.")
        if proto.get("triage_assessment") == "needs_expert_review":
            flags.append(
                f"{name}: o protocolo {proto.get('protocol_version', 'MFLAB-DF')} identificou sinais que justificam "
                "revisão pericial aprofundada, sem classificá-los isoladamente como deepfake/IA."
            )
    if not flags:
        return (
            "Nos testes automatizados de triagem executados não foram observados sinais fortes suficientes, isoladamente, "
            "para concluir pela adulteração ou geração sintética. A ausência desses sinais não prova autenticidade; "
            "técnicas adicionais, proveniência, arquivo de origem, dispositivo fonte e contexto podem ser necessários."
        )
    return "Foram observados achados que justificam exame aprofundado: " + " ".join(flags)


def _answer_question(q: str, reports: list[dict], assessment: str) -> str:
    ql = q.lower()
    if "alteração" in ql or "manipulação" in ql:
        return assessment
    if "metadad" in ql or "estrutura" in ql or "origem" in ql:
        return (
            "O laboratório registra e descreve metadados, estrutura e proveniência disponível, mas a compatibilidade "
            "com uma narrativa de origem depende de informação externa documentada e, quando possível, confronto com "
            "o arquivo/dispositivo fonte. Metadados ausentes ou editáveis não bastam, isoladamente, para confirmar ou negar a narrativa."
        )
    if "sintét" in ql or "ia" in ql or "deepfake" in ql:
        protos = []
        for r in reports:
            p = r.get("methods", {}).get("deepfake_protocol") or r.get("methods", {}).get("video_deepfake_protocol")
            if isinstance(p, dict):
                protos.append(p)
        validated = sum(int(p.get("validated_external_models", 0) or 0) for p in protos)
        review = sum(p.get("triage_assessment") == "needs_expert_review" for p in protos)
        return (
            f"Foram executados {len(protos)} protocolo(s) MFLAB-DF; {review} indicou(aram) necessidade de revisão aprofundada "
            f"e {validated} saída(s) de modelo foram explicitamente marcadas como validadas no caso. O pipeline não "
            "converte heurísticas ou um único escore em certeza. Assim, salvo convergência adicional documentada por "
            "especialista, o resultado automático permanece inconclusivo quanto à geração sintética/IA."
        )
    return "Responder em conjunto com os resultados, a conclusão preliminar e as limitações do exame, sem converter escore isolado em certeza pericial."


def _short_result(method: str, value) -> str:
    if not isinstance(value, dict):
        return str(value)[:700]
    if method in {"deepfake_protocol", "video_deepfake_protocol"}:
        compact = {k: value.get(k) for k in (
            "protocol_version", "triage_assessment", "screening_observations", "evidence_families",
            "validated_external_models", "evidentiary_conclusion",
        )}
        return json.dumps(compact, ensure_ascii=False)[:1000]
    if method == "metadata":
        if value.get("kind") == "video":
            compact = {
                "kind": "video",
                "streams": len(value.get("streams", [])),
                "format_name": value.get("format", {}).get("format_name") if isinstance(value.get("format"), dict) else None,
                "duration_s": value.get("format", {}).get("duration") if isinstance(value.get("format"), dict) else None,
            }
        else:
            compact = {k: value.get(k) for k in ("kind", "format", "mode", "size")}
            compact["exif_fields"] = len(value.get("exif", {})) if isinstance(value.get("exif"), dict) else None
        return json.dumps(compact, ensure_ascii=False)
    if method == "c2pa":
        compact = {k: value.get(k) for k in ("status", "tool", "reason", "cryptographically_validated") if value.get(k) is not None}
        if value.get("status") == "success":
            compact["manifest_present"] = True
        return json.dumps(compact, ensure_ascii=False)
    if method == "histogram":
        compact = {
            "gray_dynamic_range": value.get("gray_dynamic_range"),
            "channels": {k: {x: v.get(x) for x in ("mean", "std", "entropy_bits", "zero_bins", "clipped_black_fraction", "clipped_white_fraction")} for k, v in value.get("channels", {}).items()},
        }
        return json.dumps(compact, ensure_ascii=False)[:700]
    if method in {"noise_map", "noise_residual"}:
        keys = ("residual_std", "residual_mad", "local_std_mean", "local_std_cv", "mean", "std", "mad", "blocks")
        return json.dumps({k: value.get(k) for k in keys if k in value}, ensure_ascii=False)
    if method == "frequency":
        return json.dumps({k: value.get(k) for k in ("high_low_frequency_ratio", "spectral_peak_count", "spectral_peak_indices", "quadrant_mean_cv")}, ensure_ascii=False)
    if method == "synthetic_spectral":
        f = value.get("features", {}) if isinstance(value.get("features"), dict) else {}
        return json.dumps({"screening_flags": value.get("screening_flags", []), "high_low_frequency_ratio": f.get("high_low_frequency_ratio"), "spectral_peak_count": f.get("spectral_peak_count"), "quadrant_mean_cv": f.get("quadrant_mean_cv")}, ensure_ascii=False)
    if method == "synthetic_feature_bank":
        return json.dumps({"feature_family": value.get("feature_family"), "feature_count": value.get("feature_count"), "calibrated": value.get("calibrated")}, ensure_ascii=False)
    if method in {"synthetic_ml", "synthetic_deep"}:
        return json.dumps({k: value.get(k) for k in ("status", "model_name", "predicted_label", "score", "score_synthetic", "validated", "calibrated", "reason") if k in value}, ensure_ascii=False)
    if method == "synthetic_evidence_fusion":
        return json.dumps({k: value.get(k) for k in ("convergence_level", "family_count", "independent_families", "warning") if k in value}, ensure_ascii=False)[:800]
    if method == "face_context_consistency":
        return json.dumps({"faces_detected": value.get("faces_detected"), "faces_analyzed": value.get("faces_analyzed"), "screening_flags": value.get("screening_flags")}, ensure_ascii=False)
    if method == "resampling":
        return json.dumps({k: value.get(k) for k in ("max_nonzero_autocorrelation", "short_lag_persistence", "screening_flag", "screening_threshold", "status")}, ensure_ascii=False)
    if method == "jpeg_ghost":
        return json.dumps({"minimum_error_quality": value.get("minimum_error_quality"), "minimum_mean_abs_error": value.get("minimum_mean_abs_error")}, ensure_ascii=False)
    if method == "jpeg_quantization":
        return json.dumps({"available": value.get("available"), "table_count": value.get("table_count"), "fingerprints": value.get("fingerprints")}, ensure_ascii=False)[:700]
    if method == "prnu_screen":
        return json.dumps({k: value.get(k) for k in ("residual_std", "local_energy_mean", "local_energy_cv", "blocks", "status")}, ensure_ascii=False)
    if method == "face_artifacts":
        return json.dumps({"faces_detected": value.get("faces_detected"), "faces_analyzed": value.get("faces_analyzed"), "status": value.get("status")}, ensure_ascii=False)
    if method == "steganography_lsb":
        compact = {ch: {"p1": d.get("p1"), "entropy_bits": d.get("entropy_bits"), "pvalue": d.get("pvalue")} for ch, d in value.get("channels", {}).items()}
        return json.dumps(compact, ensure_ascii=False)
    if method == "video_timing":
        return json.dumps({k: value.get(k) for k in ("frame_count", "timestamp_count", "large_gaps", "i_frames")}, ensure_ascii=False)[:700]
    if method == "video_duplicates":
        return json.dumps({k: value.get(k) for k in ("frame_count", "duplicate_count", "adjacent_near_duplicates", "mad_threshold", "median_adjacent_mad")}, ensure_ascii=False)[:700]
    if method == "video_transition_anomalies":
        return json.dumps({k: value.get(k) for k in ("frame_count", "anomaly_count", "anomalous_transitions", "threshold", "median_transition_mad", "status")}, ensure_ascii=False)[:700]
    if method == "video_motion_discontinuities":
        return json.dumps({k: value.get(k) for k in ("frame_count", "anomaly_count", "anomalies", "robust_z_threshold", "status")}, ensure_ascii=False)[:700]
    if method in {"reference_image_difference", "reference_video_alignment"}:
        return json.dumps(value, ensure_ascii=False)[:800]
    if method == "veritas_upstream_crosscheck":
        results = value.get("results", {})
        return json.dumps({"status": value.get("status"), "features": {k: v.get("status") if isinstance(v, dict) else None for k, v in results.items()}}, ensure_ascii=False)
    return json.dumps(value, ensure_ascii=False)[:650]


def _artifact_items(report: dict) -> list[dict]:
    return list((report.get("visual_artifacts") or {}).get("items") or [])


def _append_markdown_visual_appendix(lines: list[str], reports: list[dict]) -> None:
    total = sum(len(_artifact_items(r)) for r in reports)
    lines += ["", "# APÊNDICE A — ARTEFATOS VISUAIS DAS ANÁLISES", ""]
    lines += [
        "Os artefatos abaixo são produtos derivados para inspeção e documentação. Uma região visualmente saliente "
        "não aumenta, por si só, o peso probatório do método que a originou.", ""
    ]
    if not total:
        lines += ["Nenhum artefato visual foi gerado para os métodos executados.", ""]
        return
    n = 1
    for r in reports:
        items = _artifact_items(r)
        if not items:
            continue
        lines += [f"## {Path(r.get('file', 'arquivo')).name}", ""]
        for item in items:
            rel = str(item.get("path") or "").replace("\\", "/")
            if not rel:
                continue
            lines += [
                f"**Figura A.{n} — {item.get('label', item.get('id', 'Artefato visual'))}.**",
                "",
                f"![{item.get('label', item.get('id', 'Artefato visual'))}](../{rel})",
                "",
                str(item.get("caption") or ""),
            ]
            if item.get("warning"):
                lines += [f"*Limitação:* {item['warning']}"]
            lines += [""]
            n += 1


def _append_docx_visual_appendix(doc: Document, case_dir: Path, reports: list[dict]) -> None:
    doc.add_page_break()
    doc.add_heading("APÊNDICE A — ARTEFATOS VISUAIS DAS ANÁLISES", level=1)
    doc.add_paragraph(
        "Os artefatos abaixo são produtos derivados para inspeção e documentação. Uma área visualmente saliente "
        "não constitui, por si só, prova de manipulação, geração sintética ou aumento do peso probatório do método."
    )
    n = 1
    added = 0
    for r in reports:
        items = _artifact_items(r)
        if not items:
            continue
        doc.add_heading(Path(r.get("file", "arquivo")).name, level=2)
        for item in items:
            rel = str(item.get("path") or "")
            if not rel:
                continue
            img = case_dir / Path(rel)
            if not img.exists():
                continue
            try:
                doc.add_picture(str(img), width=Cm(15.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception:
                continue
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cap.add_run(f"Figura A.{n} — {item.get('label', item.get('id', 'Artefato visual'))}.")
            run.bold = True
            if item.get("caption"):
                doc.add_paragraph(str(item["caption"]))
            if item.get("warning"):
                p = doc.add_paragraph()
                p.add_run("Limitação: ").bold = True
                p.add_run(str(item["warning"]))
            n += 1
            added += 1
    if not added:
        doc.add_paragraph("Nenhum artefato visual foi gerado para os métodos executados.")


def generate_preliminary_report(case_dir: str | Path, fmt="docx"):
    case_dir = Path(case_dir)
    case = _load_case(case_dir)
    consolidated = read_json(case_dir / "report.json", {}) or {}
    reports = consolidated.get("reports", [])
    if not reports:
        reports = [json.loads(p.read_text(encoding="utf-8")) for p in sorted((case_dir / "results").glob("*.report.json"))]
    refs = _refs_for_reports(reports)
    assessment = _assessment(reports)
    out_dir = case_dir / "final"
    out_dir.mkdir(parents=True, exist_ok=True)

    if fmt == "md":
        p = out_dir / f"{case_dir.name}_laudo_preliminar.md"
        lines = [
            f"# LAUDO TÉCNICO PERICIAL PRELIMINAR - {case['id']}", "",
            f"**Objeto:** {case['title']}", "",
            "## 1. Identificação e objeto da perícia", case["scope"], "",
            "## 2. Material recebido",
        ]
        for r in reports:
            lines.append(f"- {Path(r['file']).name} - SHA-256 `{r['sha256']}`")
        lines += [
            "", "## 3. Cadeia de custódia e preservação",
            "Os originais devem permanecer preservados; o laboratório registra SHA-256, tamanho e horário de análise. A documentação interna não substitui os registros externos de coleta, recebimento, transferência, guarda e acesso.",
            "", "## 4. Metodologia",
            "O exame combina integridade/proveniência, metadados/estrutura, compressão, sinal/ruído, possíveis manipulações locais e protocolo de mídia sintética. Métodos de triagem não são conclusivos isoladamente.", "",
            "## 5. Resultados e discussão",
        ]
        for r in reports:
            lines += [f"### {Path(r['file']).name}", f"SHA-256: `{r['sha256']}`", ""]
            ds = _deepfake_summary(r)
            if ds:
                lines += [f"**Resumo deepfake:** {ds}", ""]
            for k, v in r.get("methods", {}).items():
                lines += [f"**{METHOD_LABELS.get(k, k)}:** `{_short_result(k, v)}`", ""]
        lines += ["## 6. Conclusão preliminar", assessment, "", "## 7. Respostas aos quesitos"]
        for q in case.get("questions", []):
            lines += [f"- **Quesito:** {q}", f"  - **Resposta:** {_answer_question(q, reports, assessment)}"]
        lines += [
            "", "## 8. Limitações",
            "Este documento é preliminar. A ausência de artefatos detectáveis não prova autenticidade. Metadados podem ser alterados; recompressão e plataformas podem apagar/criar sinais; detectores de IA sofrem com mudança de domínio e geradores desconhecidos. ELA, PRNU simplificado, análise espectral e outros heurísticos devem ser corroborados por métodos validados e revisão humana.",
            "", "## 9. Referências",
        ]
        lines += [f"- {x['citation']}" for x in refs]
        _append_markdown_visual_appendix(lines, reports)
        p.write_text("\n".join(lines), encoding="utf-8")
        return str(p)

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.0)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("LAUDO TÉCNICO PERICIAL PRELIMINAR")
    run.bold = True
    run.font.size = Pt(16)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(case["id"]).bold = True
    doc.add_paragraph(f"Processo/Referência: {case.get('process_number', '')}")
    doc.add_paragraph(f"Juízo/Contratante: {case.get('court', '')}")
    doc.add_paragraph(f"Perito: {case.get('expert', {}).get('name', '')} - {case.get('expert', {}).get('qualification', '')}")

    headings = [
        "1. IDENTIFICAÇÃO E OBJETO DA PERÍCIA",
        "2. MATERIAL RECEBIDO",
        "3. CADEIA DE CUSTÓDIA E PRESERVAÇÃO",
        "4. METODOLOGIA",
        "5. RESULTADOS E DISCUSSÃO",
        "6. PROTOCOLO DE MÍDIA SINTÉTICA / DEEPFAKE",
        "7. CONCLUSÃO PRELIMINAR",
        "8. RESPOSTAS AOS QUESITOS",
        "9. LIMITAÇÕES",
        "10. REFERÊNCIAS",
    ]
    for h in headings:
        doc.add_heading(h, level=1)
        if h.startswith("1."):
            doc.add_paragraph(case["scope"])
        elif h.startswith("2."):
            for r in reports:
                doc.add_paragraph(f"{Path(r['file']).name} | SHA-256: {r['sha256']}", style="List Bullet")
        elif h.startswith("3."):
            doc.add_paragraph(
                "Os arquivos originais devem ser preservados sem alteração; o laboratório registra SHA-256, tamanho e horário de análise e executa o processamento sobre cópias de trabalho. O registro interno não substitui a documentação de recebimento, transferência, guarda e acesso ao vestígio fora do laboratório."
            )
        elif h.startswith("4."):
            doc.add_paragraph(
                "O método foi selecionado conforme a natureza de cada arquivo. O pipeline combina integridade/proveniência, metadados e estrutura, compressão, sinal/ruído, manipulação local e, quando aplicável, mídia sintética/deepfake. Métodos marcados como triagem não são determinantes isoladamente."
            )
            methods = sorted({m for r in reports for m in r.get("methods", {})})
            for m in methods:
                reg = next((r.get("method_registry", {}).get(m) for r in reports if m in r.get("method_registry", {})), {}) or {}
                suffix = " [TRIAGEM]" if reg.get("screening_only") else ""
                doc.add_paragraph(METHOD_LABELS.get(m, m) + suffix, style="List Bullet")
        elif h.startswith("5."):
            for r in reports:
                doc.add_heading(Path(r["file"]).name, level=2)
                doc.add_paragraph(f"SHA-256: {r['sha256']}")
                table = doc.add_table(rows=1, cols=2)
                table.style = "Table Grid"
                table.rows[0].cells[0].text = "Método"
                table.rows[0].cells[1].text = "Resultado resumido"
                _repeat_header_row(table.rows[0])
                _cant_split_row(table.rows[0])
                for k, v in r.get("methods", {}).items():
                    row = table.add_row()
                    _cant_split_row(row)
                    c = row.cells
                    c[0].text = METHOD_LABELS.get(k, k)
                    c[1].text = _short_result(k, v)
        elif h.startswith("6."):
            doc.add_paragraph(
                "O protocolo MFLAB-DF separa triagem heurística de evidência produzida por modelos validados. Nenhum 'AI score' é convertido automaticamente em veredito. O objetivo é registrar famílias de evidência e encaminhar casos sinalizados para revisão manual e, quando configurado, modelos aprendidos avaliados no domínio de interesse."
            )
            for r in reports:
                ds = _deepfake_summary(r)
                if ds:
                    p = doc.add_paragraph(style="List Bullet")
                    p.add_run(Path(r["file"]).name + ": ").bold = True
                    p.add_run(ds)
        elif h.startswith("7."):
            doc.add_paragraph(assessment)
        elif h.startswith("8."):
            for q in case.get("questions", []):
                p = doc.add_paragraph()
                p.add_run("Quesito: ").bold = True
                p.add_run(q)
                p = doc.add_paragraph()
                p.add_run("Resposta: ").bold = True
                p.add_run(_answer_question(q, reports, assessment))
        elif h.startswith("9."):
            doc.add_paragraph(
                "O exame automatizado é preliminar. A ausência de artefatos detectáveis não prova autenticidade. Metadados podem estar ausentes ou ser alterados; recompressão, screenshots e redes sociais podem apagar ou criar vestígios; detectores de IA podem falhar fora do domínio de treinamento. PRNU simplificado não equivale a atribuição de câmera por fingerprint/PCE. Sempre que a conclusão puder afetar decisão judicial, recomenda-se revisão independente, métodos validados e, quando possível, confronto com arquivo de origem, dispositivo fonte e exemplares conhecidos."
            )
        elif h.startswith("10."):
            for x in refs:
                doc.add_paragraph(x["citation"], style="List Bullet")

    _append_docx_visual_appendix(doc, case_dir, reports)

    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Emitido em {date.today().strftime('%d/%m/%Y')}.\n{case.get('expert', {}).get('name', '[NOME DO PERITO]')}")
    out = out_dir / f"{case_dir.name}_laudo_preliminar.docx"
    doc.save(out)
    return str(out)
